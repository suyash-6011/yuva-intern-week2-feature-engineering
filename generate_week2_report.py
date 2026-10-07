import docx
from docx.shared import Inches

doc = docx.Document()

doc.add_heading('Feature Engineering and Analysis in Automotive Pricing', 0)

doc.add_heading('1. Executive Summary & Objective', level=1)
doc.add_paragraph(
    "This report performs comprehensive feature engineering on a 500-record automotive pricing dataset. "
    "The primary goal is to derive high-value engineered features from raw variables (vehicle age, mileage, brand, price) "
    "to enhance predictive modeling capability and uncover deeper economic drivers behind vehicle depreciation."
)

doc.add_heading('2. Feature Engineering Rationale & Definitions', level=1)
doc.add_paragraph(
    "Four new features were designed and extracted from the raw dataset:\n"
    "1. Price per Mile (Price / Mileage): Captures monetary residual value relative to usage.\n"
    "2. Mileage per Year (Mileage / Age): Measures driver usage intensity.\n"
    "3. Depreciation Rate (%/Year): Standardizes value loss over time relative to brand base MSRP.\n"
    "4. Vehicle Segment Categorization: Segmented cars into 'Luxury', 'Nearly New', and 'Economy' tiers with one-hot encoding."
)

doc.add_heading('3. Correlation Analysis & Visual Findings', level=1)

images = [
    ('fig1_price_vs_age.png', 'Figure 1: Vehicle Price vs. Age Segmented by Vehicle Class'),
    ('fig2_depreciation_by_brand.png', 'Figure 2: Annual Depreciation Rate Percentage Across Brands'),
    ('fig3_usage_vs_price_per_mile.png', 'Figure 3: Usage Intensity vs. Price Per Mile Interaction'),
    ('fig4_correlation_heatmap.png', 'Figure 4: Correlation Matrix of Raw and Engineered Features')
]

for img_path, caption in images:
    doc.add_heading(caption, level=2)
    doc.add_picture(img_path, width=Inches(5.5))
    doc.add_paragraph("Analysis generated during correlation and feature interaction evaluation.")

doc.add_heading('4. Recommendations for Predictive Modeling', level=1)
doc.add_paragraph(
    "1. Depreciation Rate and Mileage per Year demonstrate stronger linear correlation with resale price than raw mileage alone.\n"
    "2. One-hot encoded vehicle segments should be utilized in tree-based and linear regression models to reduce price prediction error."
)

doc.save("Automotive_Pricing_Feature_Engineering_Report.docx")
print("Report generated: Automotive_Pricing_Feature_Engineering_Report.docx")